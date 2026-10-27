"""Gera o relatório em Excel com openpyxl (formatação, regras e gráficos)."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .kpis import KPI_LABELS

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
RED_FILL = PatternFill("solid", fgColor="F8CBAD")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _write_table(ws, df: pd.DataFrame, start_row: int = 1) -> None:
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=start_row, column=j, value=col)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, val in enumerate(row, start=1):
            if isinstance(val, pd.Timestamp):
                val = val.to_pydatetime()
            c = ws.cell(row=i, column=j, value=val)
            c.border = BORDER
            if isinstance(val, float):
                c.number_format = "0.00"
    for j, col in enumerate(df.columns, start=1):
        width = max(len(str(col)), *(len(str(v)) for v in df.iloc[:, j - 1].head(200))) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(width, 60)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(df.shape[1])}{start_row + len(df)}"


def _add_trend_chart(ws, n_rows: int, kpi_col_idx: int, title: str, anchor: str) -> None:
    chart = LineChart()
    chart.title = title
    chart.height, chart.width = 7, 16
    chart.y_axis.title = title
    chart.x_axis.number_format = "dd/mm"
    data = Reference(ws, min_col=kpi_col_idx, min_row=1, max_row=n_rows + 1, max_col=kpi_col_idx + 1)
    cats = Reference(ws, min_col=1, min_row=2, max_row=n_rows + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, anchor)


def build_excel(network_daily: pd.DataFrame, cell_kpis: pd.DataFrame, worst: pd.DataFrame, out: Path) -> Path:
    wb = Workbook()

    # 1) Resumo diário por tecnologia, com gráficos de tendência
    ws = wb.active
    ws.title = "Resumo_Diario"
    pivot = network_daily.pivot(index="date", columns="tech", values=["dcr", "ho_sr", "dl_thp_mbps"])
    pivot.columns = [f"{KPI_LABELS[k]} {t}" for k, t in pivot.columns]
    pivot = pivot.reset_index().rename(columns={"date": "Data"})
    num_cols = pivot.columns[1:]
    pivot[num_cols] = pivot[num_cols].round(2)
    _write_table(ws, pivot)
    n = len(pivot)
    _add_trend_chart(ws, n, 2, "DCR (%)", "I2")
    _add_trend_chart(ws, n, 4, "HO SR (%)", "I18")
    _add_trend_chart(ws, n, 6, "DL User Thp (Mbps)", "I34")

    # 2) Worst cells com ação sugerida
    ws2 = wb.create_sheet("Worst_Cells")
    worst_view = worst.copy()
    worst_view["kpi"] = worst_view["kpi"].map(KPI_LABELS)
    _write_table(ws2, worst_view)

    # 3) KPIs por célula no período, com destaque condicional
    ws3 = wb.create_sheet("KPIs_por_Celula")
    view = cell_kpis[["cell", "site", "city", "tech", "band", "azimuth", "tilt",
                      "rrc_sr", "dcr", "ho_sr", "dl_thp_mbps"]].round(2)
    view = view.rename(columns=KPI_LABELS)
    _write_table(ws3, view)
    last = len(view) + 1
    col = {name: get_column_letter(i + 1) for i, name in enumerate(view.columns)}
    ws3.conditional_formatting.add(f"{col['DCR (%)']}2:{col['DCR (%)']}{last}",
                                   CellIsRule(operator="greaterThan", formula=["1"], fill=RED_FILL))
    ws3.conditional_formatting.add(f"{col['HO SR (%)']}2:{col['HO SR (%)']}{last}",
                                   CellIsRule(operator="lessThan", formula=["97"], fill=RED_FILL))

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
