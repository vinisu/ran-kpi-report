"""Gera uma apresentação executiva (python-pptx + matplotlib)."""

from __future__ import annotations

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from .kpis import KPI_LABELS

NAVY = RGBColor(0x1F, 0x38, 0x64)


def _trend_png(network_daily: pd.DataFrame, kpi: str, path: Path) -> Path:
    fig, ax = plt.subplots(figsize=(6, 3.2), dpi=150)
    for tech, g in network_daily.groupby("tech"):
        ax.plot(g["date"], g[kpi], marker="o", label=tech)
    ax.set_title(KPI_LABELS[kpi])
    ax.grid(alpha=0.3)
    ax.legend()
    fig.autofmt_xdate()
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return path


def _title(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9.2), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size, p.font.bold, p.font.color.rgb = Pt(26), True, NAVY


def build_pptx(network_daily: pd.DataFrame, worst: pd.DataFrame, out: Path, tmp_dir: Path) -> Path:
    tmp_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    blank = prs.slide_layouts[6]

    # Capa
    s = prs.slides.add_slide(blank)
    _title(s, "Relatório de Performance RAN 4G/5G")
    sub = s.shapes.add_textbox(Inches(0.4), Inches(1.0), Inches(9), Inches(1)).text_frame
    start, end = network_daily["date"].min(), network_daily["date"].max()
    sub.text = f"Período: {start:%d/%m/%Y} a {end:%d/%m/%Y}  ·  Dados fictícios para demonstração"
    sub.paragraphs[0].font.size = Pt(14)

    # Tendências
    for kpis in (["dcr", "ho_sr"], ["dl_thp_mbps"]):
        s = prs.slides.add_slide(blank)
        _title(s, "Tendência diária por tecnologia")
        for i, k in enumerate(kpis):
            png = _trend_png(network_daily, k, tmp_dir / f"{k}.png")
            s.shapes.add_picture(str(png), Inches(0.3 + 4.8 * i), Inches(1.1), width=Inches(4.6))

    # Top worst cells
    s = prs.slides.add_slide(blank)
    _title(s, "Top 10 worst cells (prioridade)")
    top = worst.head(10)
    cols = ["cell", "kpi", "value", "threshold", "volume", "tilt"]
    table = s.shapes.add_table(len(top) + 1, len(cols), Inches(0.3), Inches(1.0),
                               Inches(9.4), Inches(0.3 * (len(top) + 1))).table
    for j, c in enumerate(cols):
        table.cell(0, j).text = c
    for i, r in enumerate(top.itertuples(index=False), start=1):
        vals = [r.cell, KPI_LABELS[r.kpi], f"{r.value:.2f}", f"{r.threshold:g}", f"{r.volume:,}", str(r.tilt)]
        for j, v in enumerate(vals):
            table.cell(i, j).text = v
    for row in table.rows:
        for cell in row.cells:
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out
